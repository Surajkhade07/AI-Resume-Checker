import io
import os
import re
from typing import Tuple, Optional

import pdfplumber
from docx import Document

try:
    import pypdf as PYPDF2
except ImportError:
    try:
        import PyPDF2 as PYPDF2
    except ImportError:
        PYPDF2 = None

from backend.util.file_utils import (
    FileParsingError,
    TextExtractionError,
    FileUploadError,
    log_error,
    log_warning,
    log_info,
    with_fallback
)

from backend.core.config import (
    MAX_FILE_SIZE_MB,
    MAX_FILE_SIZE_BYTES,
    SUPPORTED_MIME_TYPES
)

class FileParsingError(Exception):
    pass

class FileValidationError(Exception):
    pass

def _detect_file_type(file_data: bytes, filename: str) -> Optional[str]:
    # 1. Check file magic bytes
    if file_data.startswith(b'%PDF'):
        return 'pdf'
    if file_data.startswith(b'PK\x03\x04'):
        # ZIP-based format (docx)
        ext = os.path.splitext(filename)[1].lower()
        if ext in ('.docx', '.doc'):
            return 'docx'
        return 'docx'
    if file_data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        # Legacy binary OLE2 doc
        return 'doc'

    # 2. Check extension
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        return 'pdf'
    elif ext == '.docx':
        return 'docx'
    elif ext == '.doc':
        return 'doc'
    elif ext == '.txt':
        return 'txt'

    # 3. Optional python-magic detection
    try:
        import magic
        mime = magic.from_buffer(file_data, mime=True)
        if mime == 'application/pdf':
            return 'pdf'
        if 'wordprocessingml' in mime or mime == 'application/zip':
            return 'docx'
        if mime == 'application/msword':
            return 'doc'
        if mime.startswith('text/'):
            return 'txt'
    except Exception:
        pass

    return None

def validate_file(file_data: bytes, filename: str) -> Tuple[bool, str, Optional[str]]:
    file_size_bytes = len(file_data)
    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024 * 1024)
        return False, (
            f"File size {size_mb:.2f}MB exceeds the maximum of {MAX_FILE_SIZE_MB} MB. "
            "Please upload a smaller file or compress your resume."
        ), None

    if file_size_bytes == 0:
        return False, "Uploaded file is empty. Please check the file you have uploaded and try again.", None

    file_type = _detect_file_type(file_data, filename)
    if not file_type:
        return False, "Unsupported file format. Please upload a PDF or DOCX file.", None

    return True, '', file_type

def _extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    if PYPDF2 is None:
        return ""
    try:
        reader = PYPDF2.PdfReader(io.BytesIO(file_data))
        for page in reader.pages:
            if '/Annots' not in page:
                continue
            annots = page['/Annots']
            for annot_ref in annots:
                try:
                    annot = annot_ref.get_object() if hasattr(annot_ref, 'get_object') else annot_ref
                    if annot.get('/Subtype') != '/Link':
                        continue
                    action = annot.get('/A', {})
                    uri = action.get('/URI', '')
                    if uri and isinstance(uri, (str, bytes)):
                        if isinstance(uri, bytes):
                            uri = uri.decode('utf-8', errors='ignore')
                        uri = uri.strip()
                        if uri.startswith('http'):
                            urls.append(uri)
                except Exception:
                    pass
    except Exception:
        pass
    return '\n'.join(set(urls))

def _extract_pdf_with_pdfplumber(file_data: bytes) -> str:
    text = ''
    with pdfplumber.open(io.BytesIO(file_data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'pdfplumber extracted no text',
            user_message='No text could be extracted from the PDF.'
        )
    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()

def _extract_pdf_with_pypdf(file_data: bytes) -> str:
    if PYPDF2 is None:
        raise TextExtractionError('PyPDF2 not installed')
    text = ''
    pdf_reader = PYPDF2.PdfReader(io.BytesIO(file_data))
    for page in pdf_reader.pages:
        pages_text = page.extract_text()
        if pages_text:
            text += pages_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'PyPDF extracted no text',
            user_message='No text could be extracted from the PDF.'
        )

    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()

def extract_text_from_pdf(file_data: bytes) -> str:
    try:
        result, used_fallback = with_fallback(
            _extract_pdf_with_pdfplumber,
            _extract_pdf_with_pypdf,
            file_data,
            log_fallback=True,
        )
        return result
    except Exception as e:
        log_error(e, context="extract_text_from_pdf")
        raise FileParsingError(
            "Failed to extract text from PDF. The PDF may be scanned, image-only, or password-protected. "
            "Please ensure it contains selectable text."
        ) from e

def extract_text_from_docx(file_data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_data))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = '\n'.join(text_parts)
        if not text.strip():
            raise FileParsingError("No text could be extracted from the document.")

        return text.strip()
    except Exception as e:
        log_error(e, context='extract_text_from_docx')
        raise FileParsingError(
            "Failed to extract text from DOCX file. Please convert to PDF and try again."
        ) from e

def extract_text_from_txt(file_data: bytes) -> str:
    try:
        return file_data.decode('utf-8')
    except UnicodeDecodeError:
        return file_data.decode('latin-1', errors='ignore')

def extract_text(file_data: bytes, file_type: str) -> str:
    if file_type == 'pdf':
        return extract_text_from_pdf(file_data)
    elif file_type == 'docx':
        return extract_text_from_docx(file_data)
    elif file_type == 'txt':
        return extract_text_from_txt(file_data)
    elif file_type == 'doc':
        raise FileParsingError(
            "Legacy .doc format is not supported. Please convert your resume to .docx or .pdf."
        )
    else:
        raise FileValidationError(f"Unsupported file type: {file_type}")

def parse_resume_file(file_data: bytes, filename: str) -> Tuple[str, dict]:
    log_info(f"Parsing file: {filename}", context="parse_resume_file")

    is_valid, error_msg, file_type = validate_file(file_data, filename)
    if not is_valid:
        log_warning(f"Validation failed for file {filename}: {error_msg}", context="parse_resume_file")
        raise FileValidationError(error_msg)

    text = extract_text(file_data, file_type)
    if not text.strip():
        raise FileParsingError("Extracted resume text was empty. Please check your document.")

    metadata = {
        'filename': filename,
        'file_type': file_type,
        'file_size_bytes': len(file_data),
        'text_length': len(text),
        'success': True,
    }
    return text, metadata